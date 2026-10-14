import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


COLUMNS = [
    "Lead ID",
    "Date",
    "Name",
    "Contact No.",
    "Email Id",
    "Alt No.",
    "Model No.",
    "Category",
    "City",
    "Voice of Customer",
]


def export_excel(leads, filename="output/Leads_Report.xlsx"):

    if not leads:
        print("No leads found.")
        return

    df = pd.DataFrame(leads)

    # Ensure all columns exist
    for col in COLUMNS:
        if col not in df.columns:
            df[col] = ""

    # Keep required order
    df = df[COLUMNS]

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False)

    wb = load_workbook(filename)
    ws = wb.active

    # Header formatting
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for column in ws.columns:

        letter = get_column_letter(column[0].column)

        max_length = max(
            len(str(cell.value)) if cell.value else 0
            for cell in column
        )

        if column[0].value == "Voice of Customer":

            ws.column_dimensions[letter].width = 60

            for cell in column:
                cell.alignment = Alignment(
                    wrap_text=True,
                    vertical="top"
                )

        else:

            ws.column_dimensions[letter].width = min(max_length + 3, 35)

            for cell in column:
                cell.alignment = Alignment(vertical="top")

    wb.save(filename)

    print(f"\nExcel saved -> {filename}")