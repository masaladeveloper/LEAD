"""Import either the original extractor sheet or the 17-column sales master."""
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from crm_database import CRM_COLUMNS, STATUSES

HEADERS = {
    'dealer':'dealer', 'call id':'lead_id', 'lead id':'lead_id', 'enq date':'enquiry_date', 'date':'enquiry_date',
    'customer name':'customer_name', 'name':'customer_name', 'contactno':'contact_no', 'contact no':'contact_no',
    'email id':'email_id', 'model':'model', 'model no':'model', 'quantity':'quantity', 'alt no':'alt_no', 'city':'city',
    'voice of customer':'voice_of_customer', 'remark':'remark', 'category':'category',
    'customer responded or not':'responded', 'relevant / irrelevent deals':'relevance',
    'relevant / irrelevant deals':'relevance', 'hot / cold / warm deals':'deal_temperature',
    'win':'win', 'lost':'lost', 'status':'status', 'next follow-up':'follow_up_date',
}


def clean_header(value):
    return ' '.join(str(value or '').replace('\xa0',' ').replace(':','').replace('.','').split()).casefold()


def marked(value):
    return str(value or '').strip().casefold() not in ('', 'no', 'false', '0', 'none', 'n/a', '-')


@dataclass
class ImportResult:
    added: int = 0
    updated: int = 0
    backup: Path | None = None


def import_workbook(database, path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book.active.iter_rows(values_only=True)
        headers = [HEADERS.get(clean_header(v), '') for v in next(rows, ())]
        if 'customer_name' not in headers or 'lead_id' not in headers:
            raise ValueError('Choose a lead master or extractor workbook with customer name and Call ID / Lead ID columns.')
        records = []
        for number, values in enumerate(rows, 2):
            record = {key:value for key, value in zip(headers, values) if key}
            if not any(value is not None and str(value).strip() for value in record.values()):
                continue
            if not str(record.get('customer_name') or '').strip():
                raise ValueError(f'Row {number}: customer name is required.')
            if not str(record.get('lead_id') or '').strip() and not (record.get('contact_no') and record.get('enquiry_date')):
                raise ValueError(f'Row {number}: provide a Call ID, or both phone and enquiry date, to prevent duplicates.')
            win, lost = marked(record.pop('win', None)), marked(record.pop('lost', None))
            if win and lost:
                raise ValueError(f'Row {number}: both Win and Lost are marked.')
            if win or lost:
                record['status'] = 'Won' if win else 'Lost'
            if record.get('status'):
                statuses = {s.casefold():s for s in STATUSES}
                status = statuses.get(str(record['status']).strip().casefold())
                if not status:
                    raise ValueError(f'Row {number}: unrecognized status {record["status"]!r}.')
                record['status'] = status
            records.append((number, record))
    finally:
        book.close()
    result = ImportResult(backup=database.backup())
    with database.transaction():
        for number, record in records:
            try:
                lead_id, created = database.upsert_source_lead(record)
                # Seed sales fields for new leads. Existing CRM decisions win over an old sheet.
                if created:
                    fields = {k:v for k,v in record.items() if k in CRM_COLUMNS and v is not None and str(v).strip()}
                    database.update_lead(lead_id, fields, protect_source=False)
                    database.add_activity(lead_id, 'import', 'Imported from Excel master')
                result.added += int(created)
                result.updated += int(not created)
            except ValueError as error:
                raise ValueError(f'Row {number}: {error}') from error
    return result
