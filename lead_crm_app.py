"""Offline Windows CRM for Hyundai/IndiaMART lead follow-up."""

from __future__ import annotations

import tkinter as tk
import os
import threading
from datetime import date
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import load_workbook

from crm_database import LeadDatabase, as_iso_date
from config import MBOX_FILE
from report_generator import generate_weekly_report
from thunderbird_sync import sync_mailbox


APP_DIR = Path(__file__).resolve().parent
# Keep user data outside the packaged program directory.  A PyInstaller build
# may be upgraded or replaced, while this file remains the stable CRM record.
DATA_DIR = Path(os.environ.get("LOCALAPPDATA", APP_DIR)) / "Muneer Lead CRM"
DATABASE_PATH = DATA_DIR / "lead_crm.db"

IMPORT_HEADERS = {
    "dealer": "dealer", "call id": "lead_id", "enq date": "enquiry_date",
    "customer name": "customer_name", "contactno": "contact_no", "contact no": "contact_no",
    "email id": "email_id", "model": "model", "quantity": "quantity",
    "voice of customer": "voice_of_customer", "remark": "remark", "category": "category",
    "customer responded or not": "responded", "relevant / irrelevent deals": "relevance",
    "relevant / irrelevant deals": "relevance", "hot / cold / warm deals": "deal_temperature",
    "win": "win", "lost": "lost", "status": "status",
}

EDIT_FIELDS = (
    ("dealer", "Dealer", "text"), ("lead_id", "Call ID", "text"),
    ("enquiry_date", "Enquiry date (YYYY-MM-DD)", "text"),
    ("customer_name", "Customer name", "text"), ("contact_no", "Contact number", "text"),
    ("email_id", "Email", "text"), ("alt_no", "Alternate number", "text"),
    ("model", "Model", "text"), ("quantity", "Quantity", "text"),
    ("category", "Source category", "combo:Hot,Warm,Cold"),
    ("city", "City", "text"), ("voice_of_customer", "Voice of customer", "multiline"),
    ("remark", "Remark / latest update", "multiline"),
    ("responded", "Customer responded", "combo:Responded,Not Responded,No Answer,Callback"),
    ("relevance", "Deal relevance", "combo:Relevant,Irrelevant"),
    ("deal_temperature", "Deal temperature", "combo:Hot,Warm,Cold"),
    ("outcome", "Outcome", "combo:,Won,Lost"),
    ("status", "Status", "combo:New,In process,Follow-up,Quotation Sent,Won,Lost,Closed"),
    ("follow_up_date", "Next follow-up (YYYY-MM-DD)", "text"),
)


def clean_header(value: object) -> str:
    return " ".join(str(value or "").replace("\xa0", " ").replace(":", "").split()).casefold()


class LeadForm(tk.Toplevel):
    def __init__(self, parent: "LeadCRMApp", lead: dict | None = None):
        super().__init__(parent)
        self.parent = parent
        self.lead = lead or {}
        self.title("Edit lead" if lead else "Add lead")
        self.geometry("720x760")
        self.transient(parent)
        self.grab_set()
        self.widgets: dict[str, tk.Widget] = {}
        canvas = tk.Canvas(self, highlightthickness=0)
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=canvas.yview)
        body = ttk.Frame(canvas, padding=14)
        body.bind("<Configure>", lambda _event: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=body, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        body.columnconfigure(1, weight=1)
        for row, (field, label, kind) in enumerate(EDIT_FIELDS):
            ttk.Label(body, text=label).grid(row=row, column=0, sticky="nw", padx=(0, 12), pady=5)
            value = self.lead.get(field, "") or ""
            if kind == "multiline":
                widget = tk.Text(body, height=4, wrap="word")
                widget.insert("1.0", str(value))
            elif kind.startswith("combo:"):
                options = kind.split(":", 1)[1].split(",")
                widget = ttk.Combobox(body, values=options, state="readonly")
                widget.set(str(value))
            else:
                widget = ttk.Entry(body)
                widget.insert(0, str(value))
            widget.grid(row=row, column=1, sticky="ew", pady=5)
            self.widgets[field] = widget
        actions = ttk.Frame(body)
        actions.grid(row=len(EDIT_FIELDS), column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(actions, text="Cancel", command=self.destroy).pack(side="right", padx=4)
        ttk.Button(actions, text="Save lead", command=self.save).pack(side="right", padx=4)

    def values(self) -> dict[str, str]:
        result = {}
        for field, widget in self.widgets.items():
            if isinstance(widget, tk.Text):
                result[field] = widget.get("1.0", "end-1c").strip()
            else:
                result[field] = widget.get().strip()
        return result

    def save(self) -> None:
        values = self.values()
        if not values["customer_name"]:
            messagebox.showerror("Customer required", "Enter the customer name before saving.", parent=self)
            return
        for key, label in (("enquiry_date", "Enquiry date"), ("follow_up_date", "Next follow-up")):
            if values[key]:
                try:
                    date.fromisoformat(values[key])
                except ValueError:
                    messagebox.showerror("Invalid date", f"{label} must use YYYY-MM-DD.", parent=self)
                    return
        try:
            if self.lead:
                self.parent.database.update_lead(self.lead["id"], values)
            else:
                database_id, _ = self.parent.database.upsert_source_lead(values)
                self.parent.database.update_lead(database_id, values)
            self.parent.refresh()
            self.destroy()
        except ValueError as error:
            messagebox.showerror("Cannot save lead", str(error), parent=self)


class LeadCRMApp(tk.Tk):
    columns = (
        ("lead_id", "Call ID", 90), ("enquiry_date", "Enquiry date", 98),
        ("customer_name", "Customer", 185), ("contact_no", "Contact", 105),
        ("model", "Model", 90), ("city", "City", 115), ("category", "Category", 80),
        ("responded", "Response", 110), ("deal_temperature", "Deal", 75),
        ("status", "Status", 110), ("follow_up_date", "Follow-up", 95),
    )

    def __init__(self):
        super().__init__()
        self.title("Muneer Lead CRM")
        self.minsize(1100, 650)
        self.database = LeadDatabase(DATABASE_PATH)
        self.syncing = False
        self.protocol("WM_DELETE_WINDOW", self.close)
        self.search_text = tk.StringVar()
        self.summary_text = tk.StringVar()
        self._build()
        self.refresh()

    def _build(self) -> None:
        toolbar = ttk.Frame(self, padding=(12, 12, 12, 6))
        toolbar.pack(fill="x")
        ttk.Button(toolbar, text="Add lead", command=lambda: LeadForm(self)).pack(side="left", padx=(0, 6))
        ttk.Button(toolbar, text="Edit selected", command=self.edit_selected).pack(side="left", padx=6)
        ttk.Button(toolbar, text="Import master report", command=self.import_master).pack(side="left", padx=6)
        self.sync_button = ttk.Button(toolbar, text="Sync new emails", command=self.sync_new_emails)
        self.sync_button.pack(side="left", padx=6)
        ttk.Button(toolbar, text="Generate weekly report", command=self.generate_report).pack(side="left", padx=6)
        ttk.Button(toolbar, text="Settings", command=self.open_settings).pack(side="left", padx=6)
        ttk.Label(toolbar, text="Search:").pack(side="left", padx=(28, 5))
        search = ttk.Entry(toolbar, textvariable=self.search_text, width=28)
        search.pack(side="left")
        search.bind("<Return>", lambda _event: self.refresh())
        ttk.Button(toolbar, text="Search", command=self.refresh).pack(side="left", padx=5)
        ttk.Button(toolbar, text="Clear", command=self.clear_search).pack(side="left")

        ttk.Label(self, textvariable=self.summary_text, padding=(12, 2, 12, 8)).pack(anchor="w")
        frame = ttk.Frame(self, padding=(12, 0, 12, 12))
        frame.pack(fill="both", expand=True)
        fields = [key for key, _, _ in self.columns]
        self.tree = ttk.Treeview(frame, columns=fields, show="headings", selectmode="browse")
        for key, heading, width in self.columns:
            self.tree.heading(key, text=heading)
            self.tree.column(key, width=width, minwidth=60, anchor="w")
        y_scroll = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        x_scroll = ttk.Scrollbar(frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=y_scroll.set, xscrollcommand=x_scroll.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        y_scroll.grid(row=0, column=1, sticky="ns")
        x_scroll.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        self.tree.bind("<Double-1>", lambda _event: self.edit_selected())

    def refresh(self) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)
        for lead in self.database.search(self.search_text.get()):
            self.tree.insert("", "end", iid=str(lead["id"]), values=[lead.get(key, "") for key, _, _ in self.columns])
        counts = self.database.counts()
        self.summary_text.set(
            f"{counts['total']} leads  |  {counts['open_count']} open  |  {counts['overdue']} overdue follow-ups"
        )

    def open_settings(self) -> None:
        window = tk.Toplevel(self)
        window.title("Email sync settings")
        window.resizable(False, False)
        window.transient(self)
        window.grab_set()
        values = {
            "mailbox_path": tk.StringVar(value=self.database.get_setting("mailbox_path", str(MBOX_FILE))),
            "sender_filter": tk.StringVar(value=self.database.get_setting("sender_filter")),
            "my_email": tk.StringVar(value=self.database.get_setting("my_email", "sales.mhe@munergroup.co.in")),
        }
        labels = (
            ("mailbox_path", "Thunderbird INBOX file"),
            ("sender_filter", "Hyundai sender email / text"),
            ("my_email", "Your sales email"),
        )
        frame = ttk.Frame(window, padding=16)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)
        for row, (key, label) in enumerate(labels):
            ttk.Label(frame, text=label).grid(row=row, column=0, sticky="w", padx=(0, 10), pady=6)
            entry = ttk.Entry(frame, textvariable=values[key], width=58)
            entry.grid(row=row, column=1, sticky="ew", pady=6)
            if key == "mailbox_path":
                ttk.Button(
                    frame, text="Browse",
                    command=lambda setting_key=key: values[setting_key].set(
                        filedialog.askopenfilename(title="Select Thunderbird INBOX file") or values[setting_key].get()
                    ),
                ).grid(row=row, column=2, padx=(6, 0))

        def save() -> None:
            if not values["sender_filter"].get().strip():
                messagebox.showerror("Sender required", "Enter the Hyundai sender email/text used in the old extractor.", parent=window)
                return
            for key, variable in values.items():
                self.database.set_setting(key, variable.get().strip())
            window.destroy()

        actions = ttk.Frame(frame)
        actions.grid(row=len(labels), column=0, columnspan=3, sticky="e", pady=(12, 0))
        ttk.Button(actions, text="Cancel", command=window.destroy).pack(side="right", padx=4)
        ttk.Button(actions, text="Save settings", command=save).pack(side="right", padx=4)

    def sync_new_emails(self) -> None:
        if self.syncing:
            return
        mailbox_path = self.database.get_setting("mailbox_path", str(MBOX_FILE))
        sender_filter = self.database.get_setting("sender_filter")
        my_email = self.database.get_setting("my_email", "sales.mhe@munergroup.co.in")
        if not sender_filter:
            self.open_settings()
            return
        self.syncing = True
        self.sync_button.configure(state="disabled")
        self.summary_text.set("Syncing Thunderbird: reading only new email records. You can keep the app open.")

        def worker() -> None:
            background_db = LeadDatabase(DATABASE_PATH)
            try:
                result = sync_mailbox(background_db, mailbox_path, sender_filter, my_email)
                self.after(0, lambda: self._finish_sync(result, None))
            except Exception as error:
                self.after(0, lambda: self._finish_sync(None, error))
            finally:
                background_db.close()

        threading.Thread(target=worker, daemon=True).start()

    def _finish_sync(self, result, error: Exception | None) -> None:
        self.syncing = False
        self.sync_button.configure(state="normal")
        self.refresh()
        if error:
            messagebox.showerror("Email sync failed", str(error), parent=self)
            return
        messagebox.showinfo(
            "Email sync complete",
            f"Read {result.emails_read} new email records.\n"
            f"New leads: {result.new_leads}\nExisting leads updated: {result.existing_leads}\n"
            f"Sales-email remarks added: {result.remarks_added}",
            parent=self,
        )

    def clear_search(self) -> None:
        self.search_text.set("")
        self.refresh()

    def selected_lead(self) -> dict | None:
        selected = self.tree.selection()
        if not selected:
            messagebox.showinfo("Select a lead", "Select a lead first.", parent=self)
            return None
        return self.database.get_lead(int(selected[0]))

    def edit_selected(self) -> None:
        lead = self.selected_lead()
        if lead:
            LeadForm(self, lead)

    def import_master(self) -> None:
        path = filedialog.askopenfilename(
            title="Choose the existing IndiaMART lead master",
            filetypes=[("Excel workbook", "*.xlsx")],
        )
        if not path:
            return
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook.active
            headers = [IMPORT_HEADERS.get(clean_header(cell.value), "") for cell in next(sheet.iter_rows(max_row=1))]
            if "lead_id" not in headers or "customer_name" not in headers:
                raise ValueError("This workbook does not have the expected Call ID and Customer Name columns.")
            source_rows = []
            crm_updates = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                record = {headers[index]: value for index, value in enumerate(values) if index < len(headers) and headers[index]}
                if not record.get("lead_id") and not record.get("customer_name"):
                    continue
                record["enquiry_date"] = as_iso_date(record.get("enquiry_date"))
                record["outcome"] = "Won" if str(record.pop("win", "")).strip() else ("Lost" if str(record.pop("lost", "")).strip() else "")
                source_rows.append(record)
                crm_updates.append(record)
            added, updated = self.database.import_rows(source_rows)
            # import_rows protects existing CRM fields. For a user-selected
            # existing master, apply its weekly values intentionally afterward.
            for record in crm_updates:
                lookup = self.database.search(str(record.get("lead_id") or ""), 5)
                matched = next((item for item in lookup if str(item.get("lead_id") or "") == str(record.get("lead_id") or "")), None)
                if matched:
                    self.database.update_lead(matched["id"], record)
            self.refresh()
            messagebox.showinfo("Import complete", f"Added {added} new leads. Updated {updated} existing leads.", parent=self)
        except Exception as error:
            messagebox.showerror("Import failed", str(error), parent=self)

    def generate_report(self) -> None:
        default_name = f"India Mart Lead Master - {date.today().isoformat()}.xlsx"
        path = filedialog.asksaveasfilename(
            title="Save weekly lead report", defaultextension=".xlsx", initialfile=default_name,
            filetypes=[("Excel workbook", "*.xlsx")],
        )
        if not path:
            return
        try:
            result = generate_weekly_report(self.database.all_leads(), path)
            messagebox.showinfo("Weekly report created", f"Saved {len(self.database.all_leads())} leads to:\n{result}", parent=self)
        except Exception as error:
            messagebox.showerror("Report failed", str(error), parent=self)

    def close(self) -> None:
        self.database.close()
        self.destroy()


if __name__ == "__main__":
    LeadCRMApp().mainloop()
