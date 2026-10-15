"""HTTP and storage regressions for the PC-only app, using synthetic data."""
from io import BytesIO
from datetime import date
from unittest.mock import patch
import mailbox
import os
from pathlib import Path
import tempfile
import time
import unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from crm_database import LeadDatabase
from web_app import create_app
from thunderbird_sync import sync_mailboxes
from test_workflows import enquiry, reply


class CollectionDate(date):
    @classmethod
    def today(cls):
        return cls(2026, 10, 15)


class Web(unittest.TestCase):
    def setUp(self):
        self.clock = patch('web_app.date', CollectionDate)
        self.clock.start()
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.path = self.root / 'crm.db'
        with_db = LeadDatabase(self.path)
        self.lead_id, _ = with_db.upsert_source_lead({'lead_id':'101','customer_name':'Synthetic Buyer',
            'contact_no':'9876543210','enquiry_date':'2026-10-14','city':'Pune','model':'HX220'})
        with_db.close()
        self.client = TestClient(create_app(self.path))
        self.client.__enter__()
        self.token = {'X-CRM-Token':self.client.get('/api/session').json()['token']}

    def tearDown(self):
        self.client.__exit__(None,None,None)
        self.temp.cleanup()
        self.clock.stop()

    def request(self, method, path, body=None):
        return self.client.request(method, path, headers=self.token, json=body)

    def detail(self):
        return self.client.get(f'/api/leads/{self.lead_id}').json()

    def edit(self, values):
        d=self.detail()
        response=self.request('PATCH',f'/api/leads/{self.lead_id}',{'revision':d['lead']['revision'],'values':values})
        self.assertEqual(response.status_code,200,response.text)
        return response.json()

    def test_local_write_protection_and_host_validation(self):
        self.assertEqual(self.client.post('/api/backup').status_code,403)
        self.assertEqual(self.client.get('/api/health',headers={'host':'evil.example'}).status_code,400)
        self.assertEqual(self.request('POST','/api/backup').status_code,200)

    def test_edit_restoration_and_five_change_limit(self):
        for number in range(7):
            d=self.edit({'remark':f'Note {number}'})
        self.assertEqual(len(d['history']),5)
        self.assertEqual(d['history'][-1]['snapshot']['remark'],'Note 1')
        target=d['history'][0]
        r=self.request('POST',f'/api/leads/{self.lead_id}/history/{target["id"]}/restore',{'revision':d['lead']['revision']})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['lead']['remark'],'Note 5')
        self.assertEqual(r.json()['history'][0]['snapshot']['remark'],'Note 6')

    def test_permanent_keeps_details_notes_and_protection(self):
        self.edit({'customer_name':'Corrected','remark':'Keep this'})
        self.request('POST',f'/api/leads/{self.lead_id}/notes',{'body':'A conversation'})
        r=self.request('DELETE',f'/api/leads/{self.lead_id}/history',{'revision':self.detail()['lead']['revision']})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['history'],[])
        self.assertEqual(r.json()['lead']['remark'],'Keep this')
        self.assertEqual(r.json()['activities'][0]['body'],'A conversation')
        db=LeadDatabase(self.path)
        db.upsert_source_lead({'lead_id':'101','customer_name':'Old email'})
        self.assertEqual(db.get_lead(self.lead_id)['customer_name'],'Corrected')
        db.close()

    def test_restore_recovers_source_override_state(self):
        d=self.edit({'city':'Mumbai'})
        self.request('POST',f'/api/leads/{self.lead_id}/history/{d["history"][0]["id"]}/restore',{'revision':d['lead']['revision']})
        db=LeadDatabase(self.path)
        db.upsert_source_lead({'lead_id':'101','city':'Nagpur'})
        self.assertEqual(db.get_lead(self.lead_id)['city'],'Nagpur')
        db.close()

    def test_stale_edit_cannot_overwrite_and_noop_does_not_add_history(self):
        revision=self.detail()['lead']['revision']
        d=self.edit({'status':'Won'})
        stale=self.request('PATCH',f'/api/leads/{self.lead_id}',{'revision':revision,'values':{'status':'Lost'}})
        self.assertEqual(stale.status_code,409)
        unchanged=self.edit({'status':'Won'})
        self.assertEqual(len(unchanged['history']),len(d['history']))
        self.assertEqual(unchanged['lead']['outcome'],'Won')

    def test_validation_and_conflicts_rollback_history(self):
        db=LeadDatabase(self.path)
        db.upsert_source_lead({'lead_id':'202','customer_name':'Other'})
        db.close()
        revision=self.detail()['lead']['revision']
        for values in ({'lead_id':'202'}, {'enquiry_date':'invalid'}, {'customer_name':''}, {'quantity':0}, {'quantity':10**30}):
            r=self.request('PATCH',f'/api/leads/{self.lead_id}',{'revision':revision,'values':values})
            self.assertEqual(r.status_code,400,r.text)
        self.assertEqual(self.detail()['history'],[])
        self.assertEqual(self.detail()['lead']['lead_id'],'101')

    def test_search_across_notes_source_fields_and_pagination(self):
        self.request('POST',f'/api/leads/{self.lead_id}/notes',{'body':'special conversation'})
        for term in ('9876543210','Synthetic','101','HX220','Pune','special conversation'):
            self.assertEqual(len(self.client.get('/api/leads',params={'q':term}).json()['leads']),1)
        db=LeadDatabase(self.path)
        for n in range(31):
            db.upsert_source_lead({'lead_id':str(500+n),'customer_name':f'Buyer {n}', 'enquiry_date':'2026-01-01'})
        db.close()
        self.assertTrue(self.client.get('/api/leads').json()['has_more'])
        self.assertEqual(len(self.client.get('/api/leads?page=2').json()['leads']),7)
        self.assertEqual(self.client.get('/api/leads?size=1000').status_code,422)

    def test_schedule_complete_reminder_and_follow_up_filter(self):
        r=self.request('POST',f'/api/leads/{self.lead_id}/reminders',{'due_at':'2000-01-01T10:00','note':'Call about quotation'})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['lead']['follow_up_date'],'2000-01-01')
        reminder=r.json()['reminders'][0]
        self.assertEqual(self.request('POST',f'/api/reminders/{reminder["id"]}/complete').status_code,200)
        self.assertTrue(self.client.get('/api/reminders').json()[0]['completed_at'])
        self.assertEqual(self.detail()['lead']['follow_up_date'],'')
        self.edit({'follow_up_date':'2000-01-01'})
        self.assertEqual(len(self.client.get('/api/leads?status=Follow-ups%20due').json()['leads']),1)
        self.edit({'status':'Won'})
        self.assertEqual(self.client.get('/api/leads?status=Follow-ups%20due').json()['leads'],[])
        self.assertEqual(self.request('POST',f'/api/leads/{self.lead_id}/reminders',{'due_at':'bad'}).status_code,400)

    def test_report_uses_supplied_template_and_cumulative_cutoff(self):
        db=LeadDatabase(self.path)
        db.upsert_source_lead({'lead_id':'old','customer_name':'Older','enquiry_date':'2025-01-01'})
        db.close()
        response=self.client.get('/api/report?year=2025&through=2025-12-31')
        self.assertEqual(response.status_code,200)
        book=load_workbook(BytesIO(response.content))
        template=load_workbook('templates/lead-master.xlsx')
        self.assertEqual([c.value for c in book.active[1]], [c.value for c in template.active[1]])
        self.assertEqual(book.active['D2'].value,'Older')
        self.assertEqual(book.active.max_row,2)
        self.assertEqual(book.active.column_dimensions['F'].width,template.active.column_dimensions['F'].width)
        self.assertEqual(book.active['C2'].number_format,'dd"-"mm"-"yyyy')
        book.close();template.close()
        self.assertEqual(self.client.get('/api/report?through=bad').status_code,400)

    def test_settings_validation_and_background_sync(self):
        self.assertEqual(self.request('POST','/api/sync',{}).status_code,400)
        path=self.root/'INBOX'; box=mailbox.mbox(path)
        box.add(enquiry('303','<lead303@example.test>'));box.close();os.utime(path,(1,1))
        settings={'mailbox_paths':[str(path)],'sender_filter':'leads@example.test','my_email':'sales@example.test'}
        bad={**settings,'mailbox_paths':[str(path)+'.msf']}
        self.assertEqual(self.request('PUT','/api/settings',bad).status_code,400)
        self.assertEqual(self.request('PUT','/api/settings',settings).status_code,200)
        self.assertEqual(self.request('POST','/api/sync',{}).status_code,202)
        for _ in range(100):
            status=self.client.get('/api/sync').json()
            if not status['running']: break
            time.sleep(.02)
        self.assertFalse(status['running'])
        self.assertEqual(status['result']['new_leads'],1)
        self.assertEqual(self.client.get('/api/leads?q=303').json()['leads'][0]['lead_id'],'303')
        self.assertTrue(list((self.root/'backups').glob('*.db')))

    def test_unknown_lead_and_unsupported_fields(self):
        self.assertEqual(self.client.get('/api/leads/999').status_code,404)
        r=self.request('PATCH',f'/api/leads/{self.lead_id}',{'revision':0,'values':{'source_key':'evil'}})
        self.assertEqual(r.status_code,400)
        r=self.request('POST','/api/leads',{'revision':0,'values':{'customer_name':'Invalid','source_key':'evil'}})
        self.assertEqual(r.status_code,400)

    def test_existing_schema_is_backed_up_before_migration(self):
        import sqlite3
        legacy = self.root / 'legacy.db'
        db = LeadDatabase(legacy)
        db.upsert_source_lead({'lead_id':'legacy','customer_name':'Existing synthetic buyer'})
        db.close()
        with sqlite3.connect(legacy) as connection:
            connection.execute('ALTER TABLE leads DROP COLUMN revision')
            connection.execute('DROP TABLE lead_history')
            connection.execute('DROP TABLE reminders')
        with TestClient(create_app(legacy)) as client:
            self.assertEqual(client.get('/api/leads?all_years=true').json()['leads'][0]['customer_name'], 'Existing synthetic buyer')
        backups = list((self.root / 'backups').glob('pre-web-*.db'))
        self.assertEqual(len(backups), 1)
        with sqlite3.connect(backups[0]) as connection:
            self.assertNotIn('revision', {row[1] for row in connection.execute('PRAGMA table_info(leads)')})
            self.assertEqual(connection.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
            self.assertEqual(connection.execute('SELECT customer_name FROM leads').fetchone()[0], 'Existing synthetic buyer')

    def test_manual_lead_create_and_duplicate_validation(self):
        r=self.request('POST','/api/leads',{'revision':0,'values':{'customer_name':'Manual','lead_id':'new','status':'Follow-up'}})
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(r.json()['lead']['status'],'Follow-up')
        self.assertEqual(self.request('POST','/api/leads',{'revision':0,'values':{'customer_name':'Duplicate','lead_id':'new'}}).status_code,400)

    def wait_for_collection(self):
        for _ in range(150):
            job = self.client.get('/api/sync').json()
            if not job['running']:
                self.assertFalse(job['error'], job)
                self.assertEqual(job['result']['errors'], [], job)
                return job
            time.sleep(.02)
        self.fail('Collection did not finish')

    def prepare_range_mailbox(self):
        path = self.root / 'RANGE-INBOX'
        box = mailbox.mbox(path)
        for call_id, stamp, name in (('2501','2025-12-31','Earlier Buyer'),
                                     ('2601','2026-01-01','Alpha Buyer'),
                                     ('2615','2026-10-15','Zulu Buyer'),
                                     ('2701','2027-01-01','Future Buyer')):
            msg = enquiry(call_id, f'<{call_id}@example.test>', name)
            msg.clear_content()
            msg.set_content(f'<table><tr><td>Date</td><td>{stamp}</td></tr><tr><td>Name</td><td>{name}</td></tr><tr><td>Email</td><td>buyer@example.test</td></tr><tr><td>Voice of Customer</td><td>Needs 215L</td></tr><tr><td>Category</td><td>HOT LEAD</td></tr></table>', subtype='html')
            box.add(msg)
        box.close(); os.utime(path, (1, 1))
        r = self.request('PUT','/api/settings', {'mailbox_paths':[str(path)],'sender_filter':'leads@example.test','my_email':'sales@example.test'})
        self.assertEqual(r.status_code, 200)
        return path

    def test_year_collection_is_inclusive_saved_and_persistent(self):
        self.prepare_range_mailbox()
        self.assertEqual(self.request('POST','/api/sync',{'year':2026}).status_code,202)
        job = self.wait_for_collection()
        self.assertEqual(job['result']['new_leads'],2)
        saved = job['collection']
        self.assertEqual(saved['label'],'2026')
        self.assertEqual((saved['from_date'], saved['to_date']),('2026-01-01','2026-10-15'))
        db = LeadDatabase(self.path)
        checkpoints = dict(db.connection.execute("SELECT key,value FROM settings WHERE key LIKE 'mailbox_v2:%'"))
        db.close()
        data = self.client.get('/api/leads',params={'collection_id':saved['id']}).json()
        self.assertEqual([x['customer_name'] for x in data['leads']], ['Zulu Buyer','Synthetic Buyer','Alpha Buyer'])
        db = LeadDatabase(self.path)
        self.assertEqual(db.collections()[0]['id'],saved['id'])
        self.assertEqual(dict(db.connection.execute("SELECT key,value FROM settings WHERE key LIKE 'mailbox_v2:%'")),checkpoints)
        self.assertIsNone(db.find_by_call_id('2501'))
        self.assertIsNone(db.find_by_call_id('2701'))
        db.close()

    def test_separate_ranges_do_not_lose_skipped_older_enquiries(self):
        self.prepare_range_mailbox()
        self.request('POST','/api/sync',{'year':2026}); year26=self.wait_for_collection()['collection']
        self.request('POST','/api/sync',{'year':2025}); year25=self.wait_for_collection()['collection']
        self.assertEqual(self.client.get('/api/leads',params={'collection_id':year25['id']}).json()['total'],1)
        self.assertEqual(self.client.get('/api/leads',params={'collection_id':year26['id']}).json()['total'],3)
        self.assertEqual(len(self.client.get('/api/collections').json()),2)
        self.request('POST','/api/sync',{'collection_id':year26['id']})
        unchanged=self.wait_for_collection()
        self.assertEqual((unchanged['result']['new_leads'],unchanged['result']['updated_leads'],unchanged['result']['remarks_added']),(0,0,0))
        self.assertEqual(len(self.client.get('/api/collections').json()),2)

    def test_updates_keep_outgoing_remarks_and_request_incoming_review(self):
        path=self.prepare_range_mailbox()
        self.request('POST','/api/sync',{'year':2026}); saved=self.wait_for_collection()['collection']
        self.request('POST','/api/sync',{'year':2025}); old=self.wait_for_collection()['collection']
        db=LeadDatabase(self.path); target=db.find_by_call_id('2601'); old_lead=db.find_by_call_id('2501'); db.close()
        before=self.client.get(f'/api/leads/{target["id"]}').json()['lead']['revision']
        box=mailbox.mbox(path)
        first=reply('<sales-range@example.test>','<2601@example.test>',body='Discussed the 215L requirement.')
        second=reply('<buyer-range@example.test>','<sales-range@example.test>',sender='buyer@example.test',body='Loan arranged with a bank.')
        second.replace_header('Date','Thu, 15 Oct 2026 11:00:00 +0000')
        old_reply=reply('<old-range@example.test>','<2501@example.test>',body='Older customer update')
        for msg in (first,second,old_reply):box.add(msg)
        box.close();os.utime(path,(1,1))
        self.request('POST','/api/sync',{'collection_id':saved['id']}); job=self.wait_for_collection()
        self.assertEqual(job['result']['remarks_added'],2)
        detail=self.client.get(f'/api/leads/{target["id"]}').json()
        self.assertEqual(detail['lead']['voice_of_customer'],'Needs 215L')
        self.assertEqual(detail['lead']['category'],'Hot')
        self.assertGreater(detail['lead']['revision'],before)
        self.assertEqual(detail['remark_text'], 'Discussed the 215L requirement.')
        requests = self.client.get('/api/reply-requests').json()
        self.assertEqual(len(requests), 1)
        self.assertEqual(requests[0]['body'], 'Loan arranged with a bank.')
        self.assertEqual(requests[0]['lead_id'], target['id'])
        self.assertNotIn('Older customer',self.client.get(f'/api/leads/{old_lead["id"]}').json()['remark_text'])
        self.request('POST','/api/sync',{'collection_id':saved['id']}); repeated=self.wait_for_collection()
        self.assertEqual(repeated['result']['remarks_added'],0)
        again=self.client.get(f'/api/leads/{target["id"]}').json()
        self.assertEqual(again['lead']['revision'],detail['lead']['revision'])
        self.assertNotIn('Loan arranged', again['remark_text'])
        self.assertEqual(len(self.client.get('/api/reply-requests').json()), 1)
        self.request('POST','/api/sync',{'collection_id':old['id']}); self.wait_for_collection()
        self.assertIn('Older customer',self.client.get(f'/api/leads/{old_lead["id"]}').json()['remark_text'])
        report=self.client.get('/api/report',params={'collection_id':saved['id'],'through':'2026-10-15'})
        book=load_workbook(BytesIO(report.content)); rows=list(book.active.iter_rows(min_row=2,values_only=True))
        exported=next(row for row in rows if row[1]=='2601')
        self.assertEqual(exported[8],'Needs 215L')
        self.assertEqual(exported[9],again['remark_text'])
        self.assertNotIn('Earlier Buyer',[row[3] for row in rows]); book.close()

    def test_custom_range_and_invalid_collection_inputs(self):
        self.prepare_range_mailbox()
        for body in ({'from_date':'2026-02-01','to_date':'2026-01-01'}, {'from_date':'bad','to_date':'2026-10-15'},
                     {'year':2027}, {'collection_id':999}, {'year':2026,'from_date':'2026-01-01'},
                     {'from_date':'2026-01-01','to_date':'2027-01-01'}):
            self.assertEqual(self.request('POST','/api/sync',body).status_code,400)
        self.assertEqual(self.client.get('/api/collections').json(),[])
        self.request('POST','/api/sync',{'from_date':'2026-01-01','to_date':'2026-01-01'})
        saved=self.wait_for_collection()['collection']
        data=self.client.get('/api/leads',params={'collection_id':saved['id']}).json()
        self.assertEqual(data['total'],1)
        self.assertEqual(data['leads'][0]['customer_name'],'Alpha Buyer')
        self.assertEqual(self.client.get('/api/leads?sort=bad').status_code,400)
        self.assertEqual(self.client.get('/api/leads?collection_id=999').status_code,404)

    def test_sort_totals_page_count_and_clamping(self):
        db=LeadDatabase(self.path)
        for n in range(31):db.upsert_source_lead({'lead_id':f'sort{n}','customer_name':f'Buyer {n:02}','enquiry_date':'2026-01-01'})
        db.close()
        data=self.client.get('/api/leads').json()
        self.assertEqual((data['total'],data['pages'],data['page']),(32,2,1))
        self.assertEqual(data['leads'][0]['customer_name'],'Synthetic Buyer')
        self.assertEqual(self.client.get('/api/leads?sort=oldest').json()['leads'][0]['customer_name'],'Buyer 00')
        self.assertEqual(self.client.get('/api/leads?sort=name_asc').json()['leads'][0]['customer_name'],'Buyer 00')
        self.assertEqual(self.client.get('/api/leads?sort=name_desc').json()['leads'][0]['customer_name'],'Synthetic Buyer')
        last=self.client.get('/api/leads?page=999').json()
        self.assertEqual((last['page'],len(last['leads'])),(2,7))
        filtered=self.client.get('/api/leads?q=Buyer%200').json()
        self.assertEqual((filtered['total'],filtered['pages']),(10,1))

    def test_default_year_scopes_leads_and_report_without_collection(self):
        db = LeadDatabase(self.path)
        older, _ = db.upsert_source_lead({'lead_id':'older', 'customer_name':'Previous buyer', 'enquiry_date':'2025-05-01'})
        db.add_activity(older, 'received', 'A new reply for an older enquiry', message_id='<older-reply@test>')
        db.close()
        prefs = self.client.get('/api/preferences').json()
        self.assertEqual(prefs, {'default_enquiry_year':'current', 'resolved_year':2026, 'auto_check_mail':True})
        self.assertEqual(self.client.get('/api/leads').json()['total'], 1)
        self.assertEqual(self.client.get('/api/leads?year=2025').json()['leads'][0]['id'], older)
        self.assertEqual(self.client.get('/api/leads?all_years=true').json()['total'], 2)
        # Reply alerts include saved older leads, even when the table opens this year.
        self.assertEqual(self.client.get('/api/reply-requests').json()[0]['lead_id'], older)
        r = self.request('PUT', '/api/preferences', {'default_enquiry_year':'2025', 'auto_check_mail':False})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()['resolved_year'], 2025)
        self.assertEqual(self.client.get('/api/leads').json()['leads'][0]['id'], older)
        report = self.client.get('/api/report')
        book = load_workbook(BytesIO(report.content))
        self.assertEqual(book.active.max_row, 2)
        self.assertEqual(book.active['D2'].value, 'Previous buyer')
        book.close()
        db = LeadDatabase(self.path)
        self.assertEqual(db.get_setting('default_enquiry_year'), '2025')
        self.assertEqual(db.get_setting('auto_check_mail'), 'false')
        db.close()
        self.assertEqual(self.client.get('/api/leads?year=2025&all_years=true').status_code, 400)
        self.assertEqual(self.client.get('/api/leads?year=2027').status_code, 400)
        for invalid in ('2027', '1899', 'all', ''):
            self.assertEqual(self.request('PUT', '/api/preferences', {'default_enquiry_year':invalid}).status_code, 400)
        self.assertEqual(self.client.put('/api/preferences', json={}).status_code, 403)

    def test_reply_review_dedup_reopen_and_outgoing_cleanup_on_existing_records(self):
        db = LeadDatabase(self.path)
        db.update_lead(self.lead_id, {'remark':'Manual update'})
        db.add_activity(self.lead_id, 'sent', 'Dear Sir,\n\nCustomer needs 215L.\n\nLoan is approved.\n\nThanks & Regards,\nSales Person', message_id='<outgoing@test>')
        db.add_activity(self.lead_id, 'note', 'Internal call note')
        db.add_activity(self.lead_id, 'email', 'Legacy unknown-direction mail')
        db.add_activity(self.lead_id, 'received', 'Please call me', 'buyer@example.test', message_id='<incoming@test>')
        db.add_activity(self.lead_id, 'received', 'Please call me', 'buyer@example.test', message_id='<incoming@test>')
        db.close()
        expected = 'Manual update\n\nCustomer needs 215L.\n\nLoan is approved.'
        detail = self.detail()
        self.assertEqual(detail['remark_text'], expected)
        self.assertEqual(self.client.get('/api/leads').json()['leads'][0]['remark_preview'], expected)
        self.assertTrue(any(a['body'].startswith('Dear Sir') for a in detail['activities']))
        self.assertTrue(any(a['body'] == 'Please call me' for a in detail['activities']))
        requests = self.client.get('/api/reply-requests').json()
        self.assertEqual(len(requests), 1)
        self.assertEqual(self.request('POST', f'/api/reply-requests/{requests[0]["id"]}/review').status_code, 200)
        self.assertEqual(self.client.get('/api/reply-requests').json(), [])
        self.assertEqual(self.request('POST', '/api/reply-requests/999/review').status_code, 404)
        db = LeadDatabase(self.path)
        db.request_reply_review(requests[0]['activity_id'])
        self.assertEqual(len(db.reply_requests()), 1)
        db.close()
        self.assertEqual(self.client.get('/api/reply-requests').json()[0]['id'], requests[0]['id'])
        book = load_workbook(BytesIO(self.client.get('/api/report').content))
        self.assertEqual(book.active['J2'].value, expected)
        book.close()

    def test_automatic_check_collects_new_mail_and_stops_with_server(self):
        path = self.prepare_range_mailbox()
        # Keep the first manual collection separate from the automatic scanner.
        self.request('PUT', '/api/preferences', {'auto_check_mail':False})
        self.request('POST', '/api/sync', {'year':2026}); self.wait_for_collection()
        self.request('POST', '/api/sync', {'year':2025}); self.wait_for_collection()
        other = create_app(self.path, mail_check_seconds=.04)
        with TestClient(other) as client:
            time.sleep(.1)
            self.assertFalse(client.get('/api/sync').json()['finished_at'])
            box = mailbox.mbox(path)
            box.add(reply('<auto-in@test>', '<2601@example.test>', sender='buyer@example.test', body='Please update the quotation.'))
            box.add(reply('<old-auto-in@test>', '<2501@example.test>', sender='buyer@example.test', body='Earlier enquiry needs update.'))
            box.close(); os.utime(path, (1, 1))
            self.request('PUT', '/api/preferences', {'auto_check_mail':True})
            for _ in range(200):
                requests = client.get('/api/reply-requests').json()
                if len(requests) == 2: break
                time.sleep(.02)
            self.assertEqual(len(requests), 2)
            job = client.get('/api/sync').json()
            self.assertTrue(job['automatic'])
            self.assertEqual({r['body'] for r in requests}, {'Please update the quotation.', 'Earlier enquiry needs update.'})
            self.request('PUT', '/api/preferences', {'auto_check_mail':False})
        # Lifespan joins the checker and worker before returning.
        self.assertFalse(other.state.mail_checker.is_alive())

    def test_changed_incoming_tail_reopens_exactly_one_request(self):
        path = self.prepare_range_mailbox()
        self.request('POST', '/api/sync', {'year':2026}); self.wait_for_collection()
        item = reply('<mutable-in@test>', '<2601@example.test>', sender='buyer@example.test', body='Need a quote')
        box = mailbox.mbox(path); key = box.add(item); box.close(); os.utime(path, (1,1))
        self.request('POST', '/api/sync', {'year':2026}); self.wait_for_collection()
        request = self.client.get('/api/reply-requests').json()[0]
        self.request('POST', f'/api/reply-requests/{request["id"]}/review')
        self.request('POST', '/api/sync', {'year':2026}); self.wait_for_collection()
        self.assertEqual(self.client.get('/api/reply-requests').json(), [])
        item.set_content('Need a quote with financing options')
        box = mailbox.mbox(path); box[key] = item; box.flush(); box.close(); os.utime(path, (1,1))
        self.request('POST', '/api/sync', {'year':2026}); self.wait_for_collection()
        reopened = self.client.get('/api/reply-requests').json()
        self.assertEqual(len(reopened), 1)
        self.assertEqual(reopened[0]['id'], request['id'])
        self.assertEqual(reopened[0]['body'], 'Need a quote with financing options')
        detail = self.client.get(f'/api/leads/{request["lead_id"]}').json()
        self.assertEqual(detail['remark_text'], '')

    def test_automatic_and_manual_collection_share_worker_lock(self):
        import threading
        from thunderbird_sync import SyncResult
        self.prepare_range_mailbox()
        entered, release = threading.Event(), threading.Event()
        def slow_sync(*args, **kwargs):
            entered.set()
            release.wait(2)
            return SyncResult()
        with patch('web_app.sync_mailboxes', side_effect=slow_sync):
            other = create_app(self.path, mail_check_seconds=.03)
            with TestClient(other) as client:
                try:
                    self.assertTrue(entered.wait(1))
                    headers = {'X-CRM-Token':client.get('/api/session').json()['token']}
                    self.assertEqual(client.post('/api/sync', headers=headers, json={}).status_code, 409)
                finally:
                    release.set()




class Extraction(unittest.TestCase):
    def test_variants_plain_text_bad_quantity_and_checkpoint_recovery(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);db=LeadDatabase(root/'crm.db'); path=root/'INBOX'
            html=enquiry('201','<variant@example.test>')
            html.clear_content();html.set_content('<table><tr><td>DATE :</td><td>14-Oct-2026</td></tr><tr><td>customer name</td><td>Variant Buyer</td></tr><tr><td>mobile no:</td><td>9876543211</td></tr><tr><td>Quantity</td><td>3</td></tr></table>',subtype='html')
            plain=enquiry('202','<plain@example.test>');plain.clear_content();plain.set_content('Date: 14/10/2026\nName: Plain Buyer\nPhone: 9876543212\nModel No: HX220')
            bad=enquiry('203','<bad@example.test>');bad.clear_content();bad.set_content('<table><tr><td>Date</td><td>2026-10-14</td></tr><tr><td>Name</td><td>Bad Quantity</td></tr><tr><td>Quantity</td><td>many</td></tr></table>',subtype='html')
            box=mailbox.mbox(path)
            for msg in (html,bad,plain):box.add(msg)
            box.close();os.utime(path,(1,1))
            result=sync_mailboxes(db,[path],'leads@example.test','sales@example.test',settle_seconds=0)
            self.assertEqual(result.errors,[])
            self.assertEqual((result.new_leads,result.unparsed),(2,1))
            self.assertIn('Quantity',result.issues[0]['reason'])
            self.assertEqual(db.find_by_call_id('201')['quantity'],3)
            self.assertEqual(db.find_by_call_id('202')['customer_name'],'Plain Buyer')
            key=db.connection.execute("SELECT key FROM settings WHERE key LIKE 'mailbox_v2:%'").fetchone()[0]
            db.set_setting(key,'[invalid JSON')
            retry=sync_mailboxes(db,[path],'leads@example.test','sales@example.test',settle_seconds=0)
            self.assertEqual(retry.new_leads,0)
            self.assertEqual(db.counts()['total'],2)
            db.close()

    def test_year_cutoff_extension_revisits_previously_excluded_mail(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); db=LeadDatabase(root/'crm.db'); path=root/'INBOX'; box=mailbox.mbox(path)
            older=enquiry('401','<before-cutoff@example.test>')
            newer=enquiry('402','<after-cutoff@example.test>'); newer.clear_content()
            newer.set_content('Date: 2026-10-15\nName: Later Buyer\nPhone: 9876543212')
            box.add(older); box.add(newer); box.close();os.utime(path,(1,1))
            first=sync_mailboxes(db,[path],'leads@example.test','sales@example.test',settle_seconds=0,
                                from_date='2026-01-01',to_date='2026-10-14',scope='year:2026')
            self.assertEqual(first.new_leads,1)
            extended=sync_mailboxes(db,[path],'leads@example.test','sales@example.test',settle_seconds=0,
                                   from_date='2026-01-01',to_date='2026-10-15',scope='year:2026')
            self.assertEqual(extended.new_leads,1)
            self.assertEqual(db.counts()['total'],2)
            db.close()

    def test_changed_tail_updates_one_remark_but_header_only_changes_do_not(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); db=LeadDatabase(root/'crm.db'); path=root/'INBOX'; box=mailbox.mbox(path)
            box.add(enquiry()); item=reply(body='First response'); box.add(item); box.close();os.utime(path,(1,1))
            sync_mailboxes(db,[path],'leads@example.test','sales@example.test',settle_seconds=0,
                           from_date='2026-01-01',to_date='2026-10-15',scope='year:2026')
            box=mailbox.mbox(path); item.set_content('Updated response: loan approved'); box[1]=item;box.flush();box.close();os.utime(path,(1,1))
            changed=sync_mailboxes(db,[path],'leads@example.test','sales@example.test',settle_seconds=0,
                                   from_date='2026-01-01',to_date='2026-10-15',scope='year:2026')
            self.assertEqual(changed.remarks_added,1)
            self.assertEqual(len([a for a in db.activities(1) if a['kind']=='sent']),1)
            revision=db.get_lead(1)['revision']
            box=mailbox.mbox(path);item['X-Test-Header']='changed';box[1]=item;box.flush();box.close();os.utime(path,(1,1))
            unchanged=sync_mailboxes(db,[path],'leads@example.test','sales@example.test',settle_seconds=0,
                                     from_date='2026-01-01',to_date='2026-10-15',scope='year:2026')
            self.assertEqual(unchanged.remarks_added,0)
            self.assertEqual(db.get_lead(1)['revision'],revision)
            db.close()



if __name__ == '__main__':
    unittest.main()
