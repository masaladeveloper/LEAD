"""Regression tests with synthetic mailboxes and workbooks; no real accounts."""
import mailbox
import os
from pathlib import Path
import tempfile
import unittest
from email.message import EmailMessage

from openpyxl import Workbook, load_workbook
from crm_database import LeadDatabase
from import_service import import_workbook
from report_service import export_report
from thunderbird_sync import sync_mailboxes


def enquiry(call_id='101', message_id='<lead-101@example.test>', name='Test Customer'):
    message = EmailMessage()
    message['From'] = 'Enquiries <leads@example.test>'
    message['To'] = 'sales@example.test'
    message['Subject'] = f'Sales Inquiry - Regular - {call_id}'
    message['Message-ID'] = message_id
    message['Date'] = 'Wed, 14 Oct 2026 10:00:00 +0000'
    message.set_content('A new enquiry')
    message.add_alternative(f'<table><tr><td>Date:</td><td>2026-10-14</td></tr><tr><td>Name:</td><td>{name}</td></tr><tr><td>Contact No.:</td><td>9876543210</td></tr><tr><td>Email Id:</td><td>buyer@example.test</td></tr><tr><td>Model No.:</td><td>HX220</td></tr></table>', subtype='html')
    return message


def reply(message_id='<reply@example.test>', parent='<lead-101@example.test>', sender='sales@example.test', body='Shared the quotation.'):
    message = EmailMessage()
    message['From'] = sender
    message['To'] = 'buyer@example.test'
    message['Subject'] = 'Re: Sales Inquiry - Regular - 101'
    message['Message-ID'] = message_id
    message['Date'] = 'Wed, 14 Oct 2026 11:00:00 +0000'
    if parent:
        message['In-Reply-To'] = parent
    message.set_content(body)
    return message


class Workflows(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.database = LeadDatabase(self.root / 'leads.db')

    def tearDown(self):
        self.database.close()
        self.temp.cleanup()

    def mbox(self, name, messages, replace=False):
        path = self.root / name
        if replace and path.exists():
            path.unlink()
        box = mailbox.mbox(path)
        for message in messages:
            box.add(message)
        box.flush()
        box.close()
        os.utime(path, (1, 1))
        return str(path)

    def sync(self, paths, **kwargs):
        result = sync_mailboxes(self.database, paths, 'leads@example.test', 'sales@example.test', settle_seconds=0, **kwargs)
        self.assertEqual(result.errors, [])
        return result

    def workbook(self, rows):
        book = Workbook()
        for row in rows:
            book.active.append(row)
        path = self.root / 'master.xlsx'
        book.save(path)
        return path

    def test_repeated_sync_and_rescan_do_not_duplicate_activity(self):
        inbox = self.mbox('INBOX', [enquiry(), reply(sender='buyer@example.test')])
        sent = self.mbox('Sent', [reply('<sent@example.test>')])
        first = self.sync([sent, inbox])
        self.assertEqual((first.new_leads, first.remarks_added), (1, 2))
        second = self.sync([sent, inbox])
        self.assertEqual((second.new_leads, second.remarks_added), (0, 0))
        self.sync([inbox, sent], force=True)
        lead = self.database.all_leads()[0]
        self.assertEqual(lead['responded'], 'Responded')
        self.assertEqual(len(self.database.activities(lead['id'])), 3)

    def test_unmatched_replies_resolve_when_parent_arrives_later(self):
        sent = self.mbox('Sent', [reply(parent='<middle@example.test>'), reply('<middle@example.test>', parent='<lead-101@example.test>')])
        result = self.sync([sent])
        self.assertEqual(result.pending, 2)
        inbox = self.mbox('INBOX', [enquiry()])
        result = self.sync([inbox])
        self.assertEqual((result.pending, result.remarks_added), (0, 2))

    def test_rebuilt_mailbox_with_larger_size_rescans(self):
        path = self.mbox('INBOX', [enquiry(), enquiry('102', '<lead-102@example.test>')])
        self.sync([path])
        self.mbox('INBOX', [enquiry('103', '<lead-103@example.test>', 'Rebuilt Customer ' * 100), enquiry('104', '<lead-104@example.test>')], replace=True)
        self.assertEqual(self.sync([path]).new_leads, 2)
        self.assertEqual(self.database.counts()['total'], 4)

    def test_source_merge_and_manual_edits_are_preserved(self):
        lead_id, _ = self.database.upsert_source_lead({'lead_id':'101', 'customer_name':'First', 'city':'Pune', 'quantity':4})
        self.database.update_lead(lead_id, {'customer_name':'Corrected', 'remark':'Keep notes', 'status':'Quotation Sent'})
        self.database.upsert_source_lead({'lead_id':'101', 'customer_name':'Old email name'})
        lead = self.database.get_lead(lead_id)
        self.assertEqual((lead['customer_name'],lead['city'],lead['quantity'],lead['remark'],lead['status']), ('Corrected','Pune',4,'Keep notes','Quotation Sent'))

    def test_edited_call_id_is_not_matched_by_old_source_key(self):
        lead_id, _ = self.database.upsert_source_lead({'lead_id':'101', 'customer_name':'Original'})
        self.database.update_lead(lead_id, {'lead_id':'202'})
        new_id, created = self.database.upsert_source_lead({'lead_id':'101', 'customer_name':'New enquiry'})
        self.assertTrue(created)
        self.assertNotEqual(new_id, lead_id)
        self.assertEqual(self.database.get_lead(lead_id)['lead_id'], '202')

    def test_blank_win_lost_and_both_workbook_formats(self):
        path = self.workbook([['Call ID','Customer Name','Win','Lost','Remark'], ['101','First',None,None,'Call tomorrow']])
        result = import_workbook(self.database, path)
        self.assertTrue(result.backup.is_file())
        lead = self.database.find_by_call_id('101')
        self.assertEqual((lead['outcome'],lead['status']), ('','New'))
        self.assertEqual(lead['remark'], 'Call tomorrow')
        path = self.workbook([['Lead ID','Name','Date','Contact No.'], ['102','Second','2026-10-14',9876543210]])
        self.assertEqual(import_workbook(self.database, path).added, 1)

    def test_import_rolls_back_all_rows_on_invalid_date(self):
        path = self.workbook([['Call ID','Customer Name','Enq Date'], ['101','First','2026-10-14'], ['102','Second','not a date']])
        with self.assertRaisesRegex(ValueError, 'Row 3'):
            import_workbook(self.database, path)
        self.assertEqual(self.database.counts()['total'], 0)

    def test_reimport_does_not_replace_sales_decisions(self):
        path = self.workbook([['Call ID','Customer Name','Remark','Status'], ['101','First','Old note','New']])
        import_workbook(self.database, path)
        lead = self.database.find_by_call_id('101')
        self.database.update_lead(lead['id'], {'remark':'Current note','status':'Won'})
        import_workbook(self.database, path)
        lead = self.database.find_by_call_id('101')
        self.assertEqual((lead['remark'],lead['status'],lead['outcome']), ('Current note','Won','Won'))

    def test_no_call_ids_update_the_correct_row(self):
        path = self.workbook([['Call ID','Customer Name','ContactNo','Enq Date','Remark'], [None,'First','9876543210','2026-10-14','First note'], [None,'Second','9876543211','2026-10-14','Second note']])
        import_workbook(self.database, path)
        import_workbook(self.database, path)
        rows = self.database.all_leads()
        self.assertEqual(len(rows), 2)
        self.assertEqual([r['remark'] for r in rows], ['First note','Second note'])

    def test_export_range_activity_and_literal_formulas(self):
        lead_id, _ = self.database.upsert_source_lead({'lead_id':'101','customer_name':'=1+1','enquiry_date':'2026-10-14'})
        self.database.add_activity(lead_id, 'received', 'Please send a quote', 'buyer@example.test')
        self.database.upsert_source_lead({'lead_id':'102','customer_name':'Older','enquiry_date':'2026-09-01'})
        path = self.root / 'reports' / 'weekly.xlsx'
        self.assertEqual(export_report(self.database, path, '2026-10-01','2026-10-31'), 1)
        book = load_workbook(path)
        self.assertEqual(book.active.max_column, 17)
        self.assertEqual(book.active['D2'].data_type, 's')
        self.assertIn('Please send a quote', book.active['J2'].value)
        book.close()

    def test_status_controls_outcome_and_open_counts(self):
        lead_id, _ = self.database.upsert_source_lead({'lead_id':'101','customer_name':'First'})
        self.database.update_lead(lead_id, {'status':'Won'})
        self.assertEqual(self.database.get_lead(lead_id)['outcome'], 'Won')
        self.database.update_lead(lead_id, {'status':'Follow-up','follow_up_date':'2000-01-01'})
        self.assertEqual(self.database.get_lead(lead_id)['outcome'], '')
        self.assertEqual(self.database.counts()['open_count'], 1)
        self.assertEqual(self.database.counts()['due'], 1)

    def test_existing_database_remains_readable(self):
        self.database.connection.execute("INSERT INTO leads (source_key,lead_id,customer_name,remark,message_id) VALUES ('lead:old','old','Legacy','Preserved','<old@example.test>')")
        self.database.connection.commit()
        self.database.close()
        self.database = LeadDatabase(self.root / 'leads.db')
        self.assertEqual(self.database.find_by_call_id('old')['remark'], 'Preserved')
        self.assertEqual(self.database.find_lead_by_message_id(['<old@example.test>'])['lead_id'], 'old')

    def test_lead_without_call_id_is_promoted_when_id_arrives(self):
        original = {'customer_name':'First','contact_no':'9876543210','enquiry_date':'2026-10-14','model':'HX220'}
        lead_id, _ = self.database.upsert_source_lead(original)
        matched, created = self.database.upsert_source_lead({**original,'lead_id':'101'})
        self.assertFalse(created)
        self.assertEqual(matched, lead_id)
        self.assertEqual(self.database.get_lead(lead_id)['lead_id'], '101')

    def test_supplier_reply_is_activity_not_customer_response(self):
        path = self.mbox('INBOX', [enquiry(), reply(sender='leads@example.test')])
        result = self.sync([path])
        self.assertEqual(result.remarks_added, 1)
        self.assertEqual(self.database.all_leads()[0]['responded'], '')

    def test_supplier_reply_with_quoted_enquiry_is_not_a_new_source(self):
        message = reply(sender='leads@example.test')
        message.set_content('<p>Updated delivery is Friday.</p><blockquote><table><tr><td>Date:</td><td>2026-10-14</td></tr><tr><td>Name:</td><td>Old quoted customer</td></tr></table></blockquote>', subtype='html')
        result = self.sync([self.mbox('INBOX', [enquiry(), message])])
        self.assertEqual(result.remarks_added, 1)
        self.assertEqual(self.database.find_by_call_id('101')['customer_name'], 'Test Customer')
        self.assertEqual(self.database.activities(1)[0]['body'], 'Updated delivery is Friday.')

    def test_html_only_reply_removes_quoted_email(self):
        message = reply()
        message.set_content('<p>Current update</p><blockquote>Old quoted update</blockquote>', subtype='html')
        self.sync([self.mbox('INBOX', [enquiry(), message])])
        body = next(a['body'] for a in self.database.activities(1) if a['kind'] == 'sent')
        self.assertIn('Current update', body)
        self.assertNotIn('Old quoted', body)

    def test_blank_identifiers_cannot_create_repeat_import_duplicates(self):
        path = self.workbook([['Call ID','Customer Name'], [None,'Ambiguous customer']])
        with self.assertRaisesRegex(ValueError, 'prevent duplicates'):
            import_workbook(self.database, path)

    def test_empty_export_is_valid_and_backup_restores_data(self):
        export_report(self.database, self.root / 'empty.xlsx')
        book = load_workbook(self.root / 'empty.xlsx')
        self.assertEqual(book.active.max_column, 17)
        book.close()
        self.database.upsert_source_lead({'lead_id':'101','customer_name':'Backed up'})
        backup = self.database.backup()
        restored = LeadDatabase(backup)
        self.assertEqual(restored.find_by_call_id('101')['customer_name'], 'Backed up')
        restored.close()

    def test_exact_sender_matching(self):
        fake = enquiry()
        fake.replace_header('From', 'leads@example.test.evil.invalid')
        self.assertEqual(self.sync([self.mbox('INBOX',[fake])]).new_leads, 0)

    def test_unknown_sender_subject_does_not_attach_to_lead(self):
        inbox = self.mbox('INBOX',[enquiry(), reply(parent='', sender='unrelated@example.test')])
        result = self.sync([inbox])
        self.assertEqual((result.remarks_added,result.pending), (0,1))

    def test_missing_mailbox_is_reported_without_blocking_other_folder(self):
        inbox = self.mbox('INBOX',[enquiry()])
        result = sync_mailboxes(self.database, [str(self.root/'missing'),inbox], 'leads@example.test','sales@example.test', settle_seconds=0)
        self.assertEqual(result.new_leads, 1)
        self.assertEqual(len(result.errors), 1)


if __name__ == '__main__':
    unittest.main()
