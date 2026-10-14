"""CLI commands must operate on the same persisted master as the UI/services."""
import mailbox
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from test_workflows import enquiry


class CLI(unittest.TestCase):
    def test_import_sync_export_and_backup(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            def run(*args):
                result = subprocess.run([sys.executable, 'main.py', '--database', str(root/'crm.db'), *map(str,args)], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)
                return result.stdout
            book = Workbook()
            book.active.append(['Lead ID','Name','Date'])
            book.active.append(['200','Imported buyer','2026-10-01'])
            book.save(root/'import.xlsx')
            self.assertIn('1 added', run('import',root/'import.xlsx'))
            box = mailbox.mbox(root/'INBOX')
            box.add(enquiry())
            box.close()
            os.utime(root/'INBOX',(1,1))
            self.assertIn('1 new leads', run('sync','--mailbox',root/'INBOX','--sender','leads@example.test','--my-email','sales@example.test'))
            self.assertIn('Exported 2 leads', run('export',root/'master.xlsx'))
            book = load_workbook(root/'master.xlsx')
            self.assertEqual(book.active.max_row, 3)
            self.assertEqual(book.active.max_column, 17)
            book.close()
            self.assertTrue(Path(run('backup').strip()).is_file())
