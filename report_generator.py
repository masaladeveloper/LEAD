"""Generate the weekly Excel master report from the offline CRM."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from copy import copy
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


REPORT_COLUMNS = (
    ("Dealer ", "dealer"),
    ("Call ID", "lead_id"),
    ("Enq Date", "enquiry_date"),
    ("Customer Name ", "customer_name"),
    ("ContactNo", "contact_no"),
    ("Email Id:", "email_id"),
    ("Model", "model"),
    ("Quantity", "quantity"),
    ("\xa0\xa0Voice of Customer:", "voice_of_customer"),
    ("Remark", "remark"),
    ("Category:", "category"),
    ("Customer Responded or Not ", "responded"),
    ("Relevant / Irrelevent  deals", "relevance"),
    ("Hot / Cold / Warm Deals", "deal_temperature"),
    ("Win", "win"),
    ("Lost", "lost"),
    ("Status", "status"),
)

WIDTHS = (37, 11, 15, 34, 16, 41, 14, 13, 33, 40, 14, 36, 33, 30, 8, 8, 15)


def _date_value(value: Any) -> Any:
    if not value:
        return ""
    try:
        return datetime.strptime(str(value), "%Y-%m-%d").date()
    except ValueError:
        return value


def _report_value(lead: dict[str, Any], field: str) -> Any:
    if field == "win":
        return "Win" if lead.get("outcome") == "Won" else ""
    if field == "lost":
        return "Lost" if lead.get("outcome") == "Lost" else ""
    if field == "enquiry_date":
        return _date_value(lead.get(field))
    return lead.get(field, "") or ""


def generate_weekly_report(leads: Iterable[dict[str, Any]], destination: str | Path) -> Path:
    """Write the established 17-column IndiaMART lead-master format."""
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    template = Path(__file__).resolve().parent / 'templates' / 'lead-master.xlsx'
    workbook = load_workbook(template) if template.is_file() else Workbook()
    sheet = workbook.active
    has_template = template.is_file()
    row_styles = [copy(sheet.cell(2, column)._style) for column in range(1, 18)]
    table_style = copy(next(iter(sheet.tables.values())).tableStyleInfo) if sheet.tables else None
    for name in list(sheet.tables):
        del sheet.tables[name]
    if not has_template:
        sheet.title = "Lead Master"

    header_fill = PatternFill("solid", fgColor="FBD4B4")
    header_font = Font(bold=True)
    top_alignment = Alignment(vertical="top", wrap_text=True)
    for column, (heading, _) in enumerate(REPORT_COLUMNS, 1):
        cell = sheet.cell(1, column, heading)
        if not has_template:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.column_dimensions[cell.column_letter].width = WIDTHS[column - 1]
    if not has_template:
        sheet.row_dimensions[1].height = 33
    row_height = sheet.row_dimensions[2].height or 48

    row_number = 2
    for lead in leads:
        for column, (_, field) in enumerate(REPORT_COLUMNS, 1):
            cell = sheet.cell(row_number, column, _report_value(lead, field))
            if has_template:
                cell._style = copy(row_styles[column - 1])
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if not has_template:
                cell.alignment = top_alignment
            if field == "enquiry_date" and cell.value:
                cell.number_format = 'dd"-"mm"-"yyyy' if has_template else "d/m/yyyy"
        sheet.row_dimensions[row_number].height = row_height
        row_number += 1

    # Excel tables require at least one data row. A blank placeholder keeps an
    # empty weekly report valid and visibly ready for data.
    last_row = max(2, row_number - 1)
    table = Table(displayName="Table1" if has_template else "LeadMaster", ref=f"A1:Q{last_row}")
    table.tableStyleInfo = table_style or TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:Q{last_row}"
    workbook.save(destination)
    return destination
